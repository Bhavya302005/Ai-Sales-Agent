import csv
import hashlib
import io
import re
from datetime import UTC, datetime, timedelta
from typing import Annotated, Any
from urllib.parse import urlsplit
from uuid import UUID
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    Response,
    UploadFile,
    status,
)
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import Auth
from app.config import Settings, get_settings
from app.contact_secrets import seal_phone_number
from app.db import get_session
from app.discovery.connectors import FetchedSource, SourceCandidate
from app.discovery.ingestion import ingest_source
from app.extraction.service import extract_source
from app.notifications import notify_roles
from app.persistence.models import (
    AuditLog,
    Campaign,
    CampaignLead,
    ConsentRecord,
    Contact,
    SourceDocument,
)
from app.rate_limits import enforce_rate_limit

router = APIRouter(prefix="/api/v1/leads", tags=["lead-import"])
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 100
FIELDS = (
    "company",
    "requirement",
    "source_url",
    "contact_name",
    "phone",
    "location",
    "timezone",
    "consent_basis",
)


class LeadImportResponse(BaseModel):
    received: int
    imported: int
    duplicates: int
    rejected: int
    errors: list[str]


def _rows_from_csv(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must use UTF-8 encoding") from exc
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames or not set(FIELDS).issubset(reader.fieldnames):
        raise ValueError(f"File must contain columns: {', '.join(FIELDS)}")
    return [dict(row) for row in reader]


def _rows_from_xlsx(content: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    header = [str(value or "").strip() for value in next(values, ())]
    if not set(FIELDS).issubset(header):
        raise ValueError(f"File must contain columns: {', '.join(FIELDS)}")
    rows = [dict(zip(header, row, strict=False)) for row in values]
    workbook.close()
    return rows


def _normalize_phone(value: object) -> str:
    phone = re.sub(r"[^\d+]", "", str(value or "").strip())
    if not re.fullmatch(r"\+[1-9]\d{7,14}", phone):
        raise ValueError("phone must use E.164 format, for example +919876543210")
    return phone


def _validate_row(row: dict[str, Any], index: int) -> dict[str, str]:
    clean = {name: str(row.get(name) or "").strip() for name in FIELDS}
    if not clean["company"] or not clean["requirement"]:
        raise ValueError(f"row {index}: company and requirement are required")
    if not clean["consent_basis"]:
        raise ValueError(f"row {index}: consent_basis is required")
    clean["phone"] = _normalize_phone(clean["phone"])
    if clean["source_url"]:
        url = urlsplit(clean["source_url"])
        if url.scheme != "https" or not url.hostname:
            raise ValueError(f"row {index}: source_url must be a public HTTPS URL")
    try:
        ZoneInfo(clean["timezone"] or "Asia/Kolkata")
    except (KeyError, ValueError) as exc:
        raise ValueError(f"row {index}: timezone is not recognized") from exc
    return clean


def _configured_test_number(settings: Settings) -> str | None:
    if settings.voice_transport == "omnidim" and settings.omnidim_test_to_number:
        return settings.omnidim_test_to_number.get_secret_value()
    return (
        settings.twilio_test_to_number.get_secret_value()
        if settings.twilio_test_to_number
        else None
    )


@router.get("/export.csv")
def export_discovery_csv(
    auth: Auth,
    session: Annotated[Session, Depends(get_session)],
    opportunity_type: Annotated[str | None, Query(alias="type")] = None,
    source: str | None = None,
    actionable: bool | None = None,
) -> Response:
    statement = select(SourceDocument).where(
        SourceDocument.organization_id == auth.organization_id,
        SourceDocument.opportunity_type.is_not(None),
    )
    if opportunity_type:
        statement = statement.where(SourceDocument.opportunity_type == opportunity_type)
    if source:
        statement = statement.where(SourceDocument.source_type == source)
    if actionable is not None:
        statement = statement.where(SourceDocument.discovery_actionable == actionable)
    documents = session.scalars(statement.order_by(SourceDocument.observed_at.desc())).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "title",
            "company",
            "location",
            "opportunity_type",
            "actionable",
            "source",
            "original_url",
            "published_at",
            "evidence_excerpt",
            "rights_note",
        ]
    )
    for document in documents:
        writer.writerow(
            [
                document.discovery_title or "",
                document.discovery_company or "",
                document.discovery_location or "",
                document.opportunity_type or "weak_signal",
                "yes" if document.discovery_actionable else "no",
                document.source_type,
                document.canonical_url,
                document.published_at.isoformat() if document.published_at else "",
                document.evidence_excerpt,
                document.rights_note,
            ]
        )
    return Response(
        output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=discovery-results.csv"},
    )


@router.post("/import", response_model=LeadImportResponse)
async def import_leads(
    request: Request,
    auth: Auth,
    session: Annotated[Session, Depends(get_session)],
    settings: Annotated[Settings, Depends(get_settings)],
    file: Annotated[UploadFile, File()],
    campaign_id: Annotated[UUID, Form()],
) -> LeadImportResponse:
    if auth.role not in {"owner", "operator"}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Editor role required")
    enforce_rate_limit(
        session, identity=f"{auth.organization_id}:{auth.user_id}", category="lead-import", limit=5
    )
    campaign = session.scalar(
        select(Campaign).where(
            Campaign.organization_id == auth.organization_id,
            Campaign.id == campaign_id,
            Campaign.mode == "calling_only",
        )
    )
    if campaign is None:
        raise HTTPException(status_code=404, detail="Calling-only campaign not found")
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Lead file exceeds the 5 MB limit")
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    if suffix not in {"csv", "xlsx"}:
        raise HTTPException(status_code=422, detail="Only CSV and XLSX files are supported")
    try:
        rows = _rows_from_csv(content) if suffix == "csv" else _rows_from_xlsx(content)
    except (BadZipFile, ValueError, OSError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if not rows or len(rows) > MAX_ROWS:
        raise HTTPException(status_code=422, detail="Lead file must contain 1 to 100 rows")
    errors: list[str] = []
    validated: list[dict[str, str]] = []
    for index, row in enumerate(rows, start=2):
        try:
            validated.append(_validate_row(row, index))
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        return LeadImportResponse(
            received=len(rows), imported=0, duplicates=0, rejected=len(errors), errors=errors[:20]
        )

    imported = 0
    duplicates = 0
    for row in validated:
        row_hash = hashlib.sha256(
            "|".join(row[name].casefold() for name in FIELDS).encode()
        ).hexdigest()
        source_url = row["source_url"] or f"upload://lead/{row_hash}"
        content_text = (
            f"{row['company']} is seeking a provider and requires {row['requirement']}. "
            f"Location: {row['location'] or 'unknown'}."
        )
        ingestion = ingest_source(
            session,
            auth.organization_id,
            FetchedSource(
                candidate=SourceCandidate(
                    external_id=f"upload:{row_hash}",
                    canonical_url=source_url,
                    source_type=f"client_{suffix}",
                    rights_note=(
                        "Client-supplied lead with declared consent basis; "
                        "original source retained."
                    ),
                    published_at=None,
                    snapshot_ref=None,
                    discovery_title=row["requirement"][:500],
                    discovery_company=row["company"][:300],
                    discovery_location=row["location"][:300] or None,
                    opportunity_type="direct_requirement",
                    discovery_actionable=True,
                    provider_metadata={"import": suffix, "consent_basis": row["consent_basis"]},
                ),
                content=content_text.encode(),
            ),
        )
        if not ingestion.created:
            # Continue through contact upsert so a previously redacted import can be
            # repaired securely by re-uploading the same source row.
            duplicates += 1
        outcome = extract_source(
            session, organization_id=auth.organization_id, source=ingestion.document
        )
        lead = outcome.lead
        if lead is None or lead.company_id is None:
            errors.append(f"{row['company']}: requirement did not match the approved offering")
            continue
        phone_hash = hashlib.sha256(row["phone"].encode()).hexdigest()
        contact = session.scalar(
            select(Contact).where(
                Contact.organization_id == auth.organization_id,
                Contact.identifier_hash == phone_hash,
            )
        )
        is_test_number = row["phone"] == _configured_test_number(settings)
        encrypted_reference = (
            (
                "env:OMNIDIM_TEST_TO_NUMBER"
                if settings.voice_transport == "omnidim"
                else "env:TWILIO_TEST_TO_NUMBER"
            )
            if is_test_number
            else seal_phone_number(row["phone"], settings)
        )
        if contact is None:
            contact = Contact(
                organization_id=auth.organization_id,
                company_id=lead.company_id,
                display_name=row["contact_name"] or "Imported contact",
                channel="phone",
                identifier_encrypted_ref=encrypted_reference,
                identifier_hash=phone_hash,
                verification_status="verified" if is_test_number else "consent_attestation_required",
                demo_test_contact=is_test_number,
            )
            session.add(contact)
            session.flush()
        elif contact.identifier_encrypted_ref == "redacted:client_import":
            contact.identifier_encrypted_ref = encrypted_reference
            contact.verification_status = (
                "verified" if is_test_number else "consent_attestation_required"
            )
        consent = session.scalar(
            select(ConsentRecord).where(
                ConsentRecord.organization_id == auth.organization_id,
                ConsentRecord.contact_id == contact.id,
                ConsentRecord.channel == "browser_voice",
                ConsentRecord.status == "active",
            )
        )
        if consent is None:
            session.add(
                ConsentRecord(
                    organization_id=auth.organization_id,
                    contact_id=contact.id,
                    channel="browser_voice",
                    purpose="hackathon_demo_qualification",
                    scope="single_browser_session",
                    source=f"client_import:{row['consent_basis'][:150]}",
                    status="active",
                    recorded_at=datetime.now(UTC),
                    expires_at=datetime.now(UTC) + timedelta(days=1),
                )
            )
        linked = session.scalar(
            select(CampaignLead).where(
                CampaignLead.organization_id == auth.organization_id,
                CampaignLead.campaign_id == campaign.id,
                CampaignLead.lead_id == lead.id,
            )
        )
        if linked is None:
            session.add(
                CampaignLead(
                    organization_id=auth.organization_id,
                    campaign_id=campaign.id,
                    lead_id=lead.id,
                    approved_at=None,
                    approved_by=None,
                    owner_id=auth.user_id,
                    state="pending_review",
                )
            )
        imported += 1
        session.add(
            AuditLog(
                organization_id=auth.organization_id,
                actor_id=auth.user_id,
                action="calling_only_lead_imported",
                target_type="lead",
                target_id=lead.id,
                reason=f"client_{suffix}; phone stored as an encrypted reference plus hash",
                request_id=request.state.request_id,
            )
        )
    notify_roles(
        session,
        organization_id=auth.organization_id,
        notification_type="lead_import_completed",
        severity="success" if imported else "warning",
        title="Lead import processed",
        summary=f"Imported {imported}; duplicates {duplicates}; rejected {len(errors)}.",
        action_url="/campaigns",
        dedupe_key=f"lead-import:{request.state.request_id}",
    )
    session.commit()
    return LeadImportResponse(
        received=len(rows),
        imported=imported,
        duplicates=duplicates,
        rejected=len(errors),
        errors=errors[:20],
    )
