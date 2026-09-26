import io
import json

import httpx
import pytest
from docx import Document
from pydantic import SecretStr

from app.business_profile import (
    BusinessProfileError,
    ProfileSource,
    analyze_business_profile,
    extract_uploaded_document,
)
from app.config import Settings


def test_docx_text_is_extracted_without_persisting_the_file() -> None:
    document = Document()
    document.add_heading("Capabilities", 1)
    document.add_paragraph("We provide secure SharePoint migration and governance services.")
    content = io.BytesIO()
    document.save(content)

    extracted = extract_uploaded_document(
        "capabilities.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content.getvalue(),
    )

    assert "SharePoint migration" in extracted


def test_document_parser_rejects_malformed_or_unsupported_content() -> None:
    with pytest.raises(BusinessProfileError):
        extract_uploaded_document("payload.exe", "application/octet-stream", b"unsafe")
    with pytest.raises(BusinessProfileError):
        extract_uploaded_document("broken.pdf", "application/pdf", b"not a pdf")


def test_grounding_failure_uses_honest_fallback() -> None:
    settings = Settings(
        app_env="test",
        gemini_api_key=SecretStr("synthetic-key"),
    )

    def handler(_: httpx.Request) -> httpx.Response:
        ungrounded = {
            "company_name": "Northstar",
            "description": "A fabricated enterprise offering that is long enough.",
            "services": ["Quantum consulting"],
            "geographies": ["Mars"],
            "industries": ["Banking"],
            "customer_needs": ["Teleportation"],
            "target_customers": ["Banks"],
            "facts": {"claim": "Invented"},
            "exclusions": ["None"],
            "pricing_policy": "Ask a human for every price.",
            "qualification_questions": ["What do you need?"],
            "handoff_conditions": ["Pricing request"],
            "evidence_quotes": ["This quote is absent"],
        }
        return httpx.Response(
            200,
            json={"candidates": [{"content": {"parts": [{"text": json.dumps(ungrounded)}]}}]},
        )

    source_text = "We provide SharePoint migration services for regulated organizations."
    result = analyze_business_profile(
        settings=settings,
        company_name="Northstar",
        business_details=source_text,
        services_text="SharePoint migration",
        sources=[
            ProfileSource(
                label="User input",
                kind="user_input",
                content_hash="a" * 64,
                excerpt=source_text,
            )
        ],
        source_texts=[source_text],
        transport=httpx.MockTransport(handler),
    )

    assert result.method == "deterministic"
    assert result.warning is not None
    assert result.suggestion.services == ["SharePoint migration"]


def test_ai_analysis_requests_comprehensive_grounded_company_details() -> None:
    settings = Settings(app_env="test", gemini_api_key=SecretStr("synthetic-key"))
    captured: dict[str, object] = {}
    evidence = (
        "Northstar provides SharePoint migration and governance services for regulated "
        "financial services companies across India."
    )

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        captured.update(payload)
        grounded = {
            "company_name": "Northstar",
            "description": (
                "Northstar provides SharePoint migration and governance services for regulated "
                "financial services companies across India."
            ),
            "services": ["SharePoint migration", "SharePoint governance"],
            "geographies": ["India"],
            "industries": ["Financial services"],
            "customer_needs": ["Migrate regulated content to SharePoint"],
            "target_customers": ["Regulated financial services companies"],
            "facts": {
                "Platform expertise": "SharePoint",
                "Regulatory focus": "Regulated financial services",
            },
            "exclusions": ["Organizations outside the supported evidence require review"],
            "pricing_policy": "Pricing is not specified; route commercial questions to a human.",
            "qualification_questions": ["What content and governance scope is required?"],
            "handoff_conditions": ["The prospect requests pricing or a custom commitment."],
            "evidence_quotes": [evidence],
        }
        return httpx.Response(
            200,
            json={
                "candidates": [{"content": {"parts": [{"text": json.dumps(grounded)}]}}],
                "usageMetadata": {"promptTokenCount": 100, "candidatesTokenCount": 200},
            },
        )

    result = analyze_business_profile(
        settings=settings,
        company_name="Northstar",
        business_details=evidence,
        services_text="SharePoint migration; SharePoint governance",
        sources=[
            ProfileSource(
                label="User input",
                kind="user_input",
                content_hash="b" * 64,
                excerpt=evidence,
            )
        ],
        source_texts=[evidence],
        transport=httpx.MockTransport(handler),
    )

    prompt = captured["contents"][0]["parts"][0]["text"]  # type: ignore[index]
    generation = captured["generationConfig"]  # type: ignore[assignment]
    assert result.method == "gemini"
    assert "REQUIRED COVERAGE CHECKLIST" in prompt
    assert "technology stack and platform expertise" in prompt
    assert "target buyer roles" in prompt
    assert "case studies or named customers" in prompt
    assert generation["maxOutputTokens"] == 6000  # type: ignore[index]
